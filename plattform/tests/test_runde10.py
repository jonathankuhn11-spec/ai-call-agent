"""Runde 10: direkt eingebundene Open-Source-Bausteine (Leaflet, uPlot, openpyxl, Jinja2) und Produktisierung
(Editionen, Lizenzschlüssel, Testphase, Feature-Gates, Selbstregistrierung, Betreiberkonsole, White-Label, Version)."""
from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from core import produkt, vorlagen
from core.demo import DEMO_ZEIT
from core.modelle import Mandant, OutboxEintrag

HEUTE = DEMO_ZEIT.date()


def test_ui_bibliotheken_und_pwa_werden_ausgeliefert(umgebung):
    c, *_ = umgebung
    for pfad, inhalt in (("/app/lib/leaflet.js", "Leaflet 1.9.4"), ("/app/lib/uplot.min.js", "uPlot (v1.6.31)"), ("/app/lib/leaflet.css", ".leaflet-"),
                         ("/app/manifest.webmanifest", '"display": "standalone"'), ("/app/sw.js", "shk-ui-v"), ("/app/icon.svg", "<svg")):
        r = c.get(pfad)
        assert r.status_code == 200 and inhalt in r.text, pfad
    html = c.get("/app/").text
    assert 'lib/leaflet.js' in html and 'lib/uplot.min.js' in html and 'manifest.webmanifest' in html and 'data-ansicht="betreiber"' in html
    assert c.get("/health").json()["version"] == produkt.__version__ and produkt.__version__.count(".") == 2
    p = c.get("/produkt").json()
    assert p["name"] == "SHK-Backoffice" and set(p["editionen"]) == {"starter", "pro", "enterprise"} and p["version"] == produkt.__version__


def test_excel_exporte_sind_lesbare_arbeitsmappen(umgebung):
    c, db, ids, h = umgebung
    r = c.get("/auswertungen/export.xlsx", headers=h)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats") and "auswertungen-2026-11-17.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Monate", "Monteure", "Kennzahlen"]
    assert [z.value for z in wb["Monate"][1]] == ["Monat", "Umsatz netto", "Zahlungen", "Material (Belege)"]
    assert wb["Monate"].max_row == 13 and wb["Monate"]["B2"].number_format.endswith('"€"')
    k = load_workbook(BytesIO(c.get("/kunden/export.xlsx", headers=h).content))["Kunden"]
    assert k["A2"].value == "Familie Beispiel" and k["G2"].value == "wartungsvertrag" and k["A1"].font.bold
    op = load_workbook(BytesIO(c.get("/offene-posten/export.xlsx", headers=h).content))["Offene Posten"]
    assert op["A1"].value == "Rechnung"
    t = c.get("/termine/export.xlsx?von=2026-11-01&bis=2026-12-31", headers=h)
    assert t.status_code == 200 and load_workbook(BytesIO(t.content))["Termine"]["A1"].value == "Datum"


def test_textvorlagen_standard_anpassung_sandbox_und_versand(umgebung):
    c, db, ids, h = umgebung
    v = c.get("/betrieb/vorlagen", headers=h).json()["vorlagen"]
    assert {x["schluessel"] for x in v} >= {"terminbestaetigung_email", "terminbestaetigung_sms", "terminbestaetigung_whatsapp", "rechnung_email", "angebot_email"}
    assert not any(x["angepasst"] for x in v)
    # Vorschau mit Beispieldaten, Fehler in der Vorlage wird gemeldet, Sandbox blockt Attributzugriffe
    vs = c.post("/betrieb/vorlagen/vorschau", headers=h, json={"schluessel": "terminbestaetigung_email"}).json()
    assert "Familie Beispiel" in vs["text"] and "Dienstag, 17.11.2026" in vs["text"] and vs["betreff"] == "Terminbestätigung 17.11.2026 10:00"
    assert c.post("/betrieb/vorlagen/vorschau", headers=h, json={"schluessel": "terminbestaetigung_sms", "text": "{% if %}"}).status_code == 422
    assert c.post("/betrieb/vorlagen/vorschau", headers=h, json={"schluessel": "terminbestaetigung_sms", "text": "{{ kunde.__class__ }}"}).status_code == 422
    assert c.post("/betrieb/vorlagen/vorschau", headers=h, json={"schluessel": "unbekannt"}).status_code == 422
    # Anpassung speichern und im echten Versand wirksam
    r = c.put("/betrieb/vorlagen", headers=h, json={"vorlagen": {"terminbestaetigung_email": {"betreff": "Ihr Termin bei {{ betrieb.name }}", "text": "Hallo {{ kunde.name }}, {{ termin.datum }} {{ termin.uhrzeit }} Uhr – {{ anliegen }}. {{ betrieb.telefon }}"}}})
    assert r.status_code == 200 and next(x for x in r.json()["vorlagen"] if x["schluessel"] == "terminbestaetigung_email")["angepasst"]
    assert c.put("/betrieb/vorlagen", headers=h, json={"vorlagen": {"terminbestaetigung_email": {"text": "{{ kunde.name"}}}).status_code == 422
    with db.sitzungen() as s:
        m = s.get(Mandant, ids["mandant"])
        assert vorlagen.quelle(m, "terminbestaetigung_email")["betreff"] == "Ihr Termin bei {{ betrieb.name }}"
    gesendet = []
    c.put("/integrationen/endkunde", headers=h, json={"kanaele": ["email"], "email": {"absender": "info@beispiel.de", "host": "smtp.example.test"}})
    c.patch(f"/kunden/{ids['kunde']}", headers=h, json={"email": "familie@example.test"})
    f = c.post("/servicefaelle", headers=h, json={"kunde_id": ids["kunde"], "anliegen": "Wartung Gastherme", "anlage_id": ids["anlage"]}).json()
    t = c.post(f"/servicefaelle/{f['id']}/termine", headers=h, json={"ressource_id": ids["monteur_gas"], "beginn": "2026-11-18T13:00:00+01:00"})
    assert t.status_code in (200, 201), t.text
    t = t.json()
    fr = [x for x in c.get("/freigaben", headers=h).json() if x["aktion"] == "termin_bestaetigen"][-1]
    c.post(f"/freigaben/{fr['id']}/entscheidung", headers=h, json={"freigeben": True, "wer": "buero"})
    from core.adapter.email import EmailAdapter
    with db.sitzungen() as s:
        e = s.scalars(select(OutboxEintrag).where(OutboxEintrag.aktion == "terminbestaetigung", OutboxEintrag.ziel == "endkunde").order_by(OutboxEintrag.id.desc())).first()
        assert e is not None and (t.get("id") or t.get("termin", {}).get("id") or "") in str(e.payload)
        a = EmailAdapter({"absender": "info@beispiel.de", "host": "smtp.example.test"}, versand=gesendet.append)
        a.zustellen(e, s)
    assert gesendet and gesendet[0]["Subject"] == "Ihr Termin bei Demo Haustechnik GmbH" and gesendet[0].get_content().startswith("Hallo Familie Beispiel, 18.11.2026 13:00 Uhr – Wartung Gastherme.")
    # Zurücksetzen: leerer Text löscht die Anpassung
    r = c.put("/betrieb/vorlagen", headers=h, json={"vorlagen": {"terminbestaetigung_email": {"text": "", "betreff": ""}}}).json()
    assert not next(x for x in r["vorlagen"] if x["schluessel"] == "terminbestaetigung_email")["angepasst"]
    # fehlerhafte eigene Vorlage fällt beim Versand still auf den Standard zurück
    with db.sitzungen() as s:
        m = s.get(Mandant, ids["mandant"])
        m.stammdaten = {**m.stammdaten, "vorlagen": {"terminbestaetigung_sms": {"text": "{{ nicht_da.x }}"}}}
        assert "ist bestätigt" in vorlagen.rendern(m, "terminbestaetigung_sms", {**vorlagen.BEISPIEL})["text"]


def test_lizenzschluessel_editionen_testphase_und_feature_gates(umgebung, monkeypatch):
    c, db, ids, h = umgebung
    st = c.get("/betrieb/lizenz", headers=h).json()
    assert st["quelle"] == "testphase" and st["edition"] == "enterprise" and "website" in st["features"] and st["hinweis"].startswith("Testphase")
    # Schlüsselpaar, signierter Schlüssel, Prüfung (Signatur, Mandantenbindung, Ablauf)
    paar = produkt.schluesselpaar()
    schl = produkt.schluessel_erzeugen(paar["privat"], edition="pro", kunde="Demo Haustechnik GmbH", gueltig_bis=HEUTE + timedelta(days=365))
    assert c.put("/betrieb/lizenz", headers=h, json={"schluessel": schl}).status_code == 422  # kein öffentlicher Schlüssel hinterlegt
    monkeypatch.setenv("LIZENZ_OEFFENTLICHER_SCHLUESSEL", paar["oeffentlich"])
    assert c.put("/betrieb/lizenz", headers=h, json={"schluessel": schl[:-4] + "AAAA"}).status_code == 422
    fremd = produkt.schluessel_erzeugen(paar["privat"], edition="pro", kunde="X", gueltig_bis=HEUTE + timedelta(days=30), mandant="anderer")
    assert "anderen Mandanten" in c.put("/betrieb/lizenz", headers=h, json={"schluessel": fremd}).json()["detail"]
    alt = produkt.schluessel_erzeugen(paar["privat"], edition="pro", kunde="X", gueltig_bis=HEUTE - timedelta(days=1))
    assert "abgelaufen" in c.put("/betrieb/lizenz", headers=h, json={"schluessel": alt}).json()["detail"]
    r = c.put("/betrieb/lizenz", headers=h, json={"schluessel": schl}).json()
    assert r["edition"] == "pro" and r["quelle"] == "lizenz" and r["lizenz_kunde"] == "Demo Haustechnik GmbH" and "mehrbetrieb" not in r["features"]
    # Pro darf Website und WhatsApp; Starter nicht (402), Limits greifen
    assert c.put("/betrieb/website", headers=h, json={"aktiv": True, "claim": "Test"}).status_code == 200
    assert c.put("/betreiber/mandanten/{}/edition".format(ids["mandant"]), json={"edition": "starter"}).json()["quelle"] == "betreiber"
    assert c.put("/betrieb/website", headers=h, json={"aktiv": True}).status_code == 402
    assert c.get("/dienstplan", headers=h).status_code == 402 and c.get("/upsell/uebersicht", headers=h).status_code == 402
    r = c.put("/integrationen/endkunde", headers=h, json={"kanaele": ["email", "whatsapp"], "email": {"absender": "a@b.de", "host": "h"}, "whatsapp": {"phone_number_id": "1", "token": "t"}})
    assert r.status_code == 402 and "WhatsApp" in r.json()["detail"]
    assert c.get("/auswertungen/export.xlsx", headers=h).status_code == 200  # Auswertungen sind in Starter enthalten
    c.post("/ressourcen", headers=h, json={"name": "Monteur C", "qualifikationen": []})
    r = c.post("/ressourcen", headers=h, json={"name": "Monteur D", "qualifikationen": []})
    assert r.status_code == 402 and "höchstens 3 mitarbeiter" in r.json()["detail"]
    # Betreiber: Zuweisung aufheben → zurück in die Testphase; Testphase verlängern; abgelaufene Testphase → Starter
    assert c.put(f"/betreiber/mandanten/{ids['mandant']}/edition", json={"edition": None}).json()["quelle"] == "testphase"
    bis = c.post(f"/betreiber/mandanten/{ids['mandant']}/testphase", json={"tage": 14}).json()["testphase_bis"]
    assert date.fromisoformat(bis) >= HEUTE + timedelta(days=14)
    with db.sitzungen() as s:
        m = s.get(Mandant, ids["mandant"])
        m.testphase_bis = HEUTE - timedelta(days=1)
        s.commit()
    st = c.get("/betrieb/lizenz", headers=h).json()
    assert st["quelle"] == "starter" and st["edition"] == "starter"
    assert c.put(f"/betreiber/mandanten/{ids['mandant']}/edition", json={"edition": "gold"}).status_code == 422
    assert c.put("/betreiber/mandanten/unbekannt/edition", json={"edition": "pro"}).status_code == 404


def test_betreiberkonsole_und_lizenzausstellung(umgebung, monkeypatch):
    c, db, ids, h = umgebung
    b = c.get("/betreiber/mandanten").json()
    assert b["summen"]["mandanten"] >= 1 and b["mandanten"][0]["kunden"] >= 1 and b["mandanten"][0]["quelle"] == "testphase"
    assert c.post("/betreiber/lizenzschluessel", json={"edition": "pro", "kunde": "Muster GmbH", "gueltig_bis": "2027-12-31"}).status_code == 409
    paar = produkt.schluesselpaar()
    monkeypatch.setenv("LIZENZ_PRIVATER_SCHLUESSEL", paar["privat"])
    r = c.post("/betreiber/lizenzschluessel", json={"edition": "enterprise", "kunde": "Muster GmbH", "gueltig_bis": "2027-12-31", "mandant": ids["mandant"], "sitze": 5}).json()
    nutz = produkt.schluessel_pruefen(r["schluessel"], paar["oeffentlich"], ids["mandant"], HEUTE)
    assert nutz["edition"] == "enterprise" and nutz["sitze"] == 5 and nutz["kunde"] == "Muster GmbH"
    assert c.post("/betreiber/lizenzschluessel", json={"edition": "gold", "kunde": "Muster GmbH", "gueltig_bis": "2027-12-31"}).status_code == 422
    uebersicht = produkt.mandanten_uebersicht(db.sitzungen().__enter__(), HEUTE)
    assert produkt.summen(uebersicht)["je_quelle"]["testphase"] >= 1


def test_selbstregistrierung_legt_betrieb_mit_testphase_an(umgebung, monkeypatch):
    c, db, ids, h = umgebung
    assert c.get("/registrieren").status_code == 404 and c.post("/registrierung", json={"name": "X"}).status_code == 404
    monkeypatch.setenv("REGISTRIERUNG_OFFEN", "1")
    seite = c.get("/registrieren")
    assert seite.status_code == 200 and 'name="website"' in seite.text and "30 Tage kostenlos" in seite.text and "Enterprise" in seite.text
    assert c.post("/registrierung", json={"name": "Neu", "email": "x", "datenschutz": True}).status_code == 422
    assert c.post("/registrierung", json={"name": "Neuer Betrieb", "email": "neu@beispiel.de", "datenschutz": False}).status_code == 422
    bot = c.post("/registrierung", json={"name": "Bot GmbH", "email": "bot@beispiel.de", "datenschutz": True, "website": "http://spam"}).json()
    assert bot["mandant_id"] is None
    r = c.post("/registrierung", json={"name": "Autohaus Neu", "branche": "kfz", "email": "neu@beispiel.de", "telefon": "0221 1", "ort": "Köln", "inhaber": "N. Neu", "datenschutz": True})
    assert r.status_code == 201, r.text
    a = r.json()
    assert a["edition"] == "enterprise" and a["testphase_bis"] == (HEUTE + timedelta(days=30)).isoformat() and a["app_url"].endswith("/app/#einrichtung")
    with db.sitzungen() as s:
        m = s.get(Mandant, a["mandant_id"])
        assert m.gewerk.lower() == "kfz" and m.stammdaten["email"] == "neu@beispiel.de" and m.stammdaten.get("buchung_token") and m.testphase_bis == HEUTE + timedelta(days=30)
    assert "bereits registriert" in c.post("/registrierung", json={"name": "Autohaus Neu", "branche": "kfz", "email": "neu@beispiel.de", "datenschutz": True}).json()["detail"]
    for i in range(4):
        c.post("/registrierung", json={"name": f"Betrieb {i}", "email": f"b{i}@beispiel.de", "datenschutz": True})
    assert c.post("/registrierung", json={"name": "Betrieb 9", "email": "b9@beispiel.de", "datenschutz": True}).status_code == 429
    assert len(c.get("/betreiber/mandanten").json()["mandanten"]) == 6


def test_white_label_aus_umgebung(monkeypatch):
    monkeypatch.setenv("PRODUKT_NAME", "KUHN Backoffice")
    monkeypatch.setenv("PRODUKT_FARBE", "#123456")
    monkeypatch.setenv("PRODUKT_ANBIETER", "KUHN")
    from core.demo import demo_client
    c, _ = demo_client()
    p = c.get("/produkt").json()
    assert p["name"] == "KUHN Backoffice" and p["farbe"] == "#123456" and p["anbieter"] == "KUHN" and p["registrierung_offen"] is False
    assert c.get("/openapi.json").json()["info"]["title"] == "KUHN Backoffice Core-Service"
